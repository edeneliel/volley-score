from docx import Document
from docx.oxml.ns import qn
from docx.shared import Inches
from docx.oxml import OxmlElement
from docx.enum.section import WD_ORIENTATION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT

from utilities.config import Constants


class ScoreSheet:
    def __init__(self):
        self.document = Document()

        section = self.document.sections[-1]
        section.orientation = WD_ORIENTATION.LANDSCAPE
        section.top_margin = Inches(0)
        section.left_margin = Inches(0)
        section.right_margin = Inches(0)
        section.bottom_margin = Inches(0)
        section.page_width = Inches(16.5)
        section.page_height = Inches(11.7)

    @classmethod
    def __merge_cells(cls, table, source_row, source_column, target_row, target_column):
        source = table.cell(source_row, source_column)
        target = table.cell(target_row, target_column)
        source.merge(target)
        return source

    @classmethod
    def __merge_row_cells(cls, table, row, start_column, config):
        current_column = start_column
        for header in config:
            next_column = current_column + header
            cls.__merge_cells(table, row, current_column, row, next_column - 1)
            current_column = next_column

    @classmethod
    def __set_table_set_header(cls, table, start_column):
        cls.__merge_row_cells(table=table,
                              row=0,
                              start_column=start_column,
                              config=Constants.SET_HEADERS_ORDER)

    @classmethod
    def __set_table_players_cells(cls, table, start_column, start_row):
        config = [Constants.PLAYER_CELL_WIDTH] * Constants.PLAYERS_CELLS_AMOUNT

        for current_row in range(start_row, start_row + Constants.PLAYERS_ROWS_AMOUNT):
            cls.__merge_row_cells(table=table,
                                  row=current_row,
                                  config=config,
                                  start_column=start_column)

    @classmethod
    def __set_score_columns(cls, table, start_column):
        for current_column in range(start_column, start_column + 4):
            cls.__merge_cells(table, 1, current_column, -4, current_column)

    @classmethod
    def __set_timeout_cells(cls, table, start_column):
        for current_row in [-3, -2, -1]:
            cls.__merge_cells(table, current_row, start_column, current_row, start_column + 3)

    def add_set_table(self):
        table = self.document.add_table(rows=10, cols=33)
        table.style = 'Table Grid'

        self.__merge_cells(table, 0, 0, -1, 0)
        self.__set_table_set_header(table=table, start_column=1)
        self.__set_table_players_cells(table=table, start_column=1, start_row=1)
        self.__set_table_players_cells(table=table, start_column=17, start_row=1)
        self.__set_score_columns(table=table, start_column=13)
        self.__set_score_columns(table=table, start_column=29)
        self.__set_timeout_cells(table=table, start_column=13)
        self.__set_timeout_cells(table=table, start_column=29)

        # tc = cell._tc
        # tcPr = tc.get_or_add_tcPr()
        # textDirection = OxmlElement('w:textDirection')
        # textDirection.set(qn('w:val'), 'btLr')
        # tcPr.append(textDirection)

        return table

    def save(self):
        self.document.save('test.docx')
